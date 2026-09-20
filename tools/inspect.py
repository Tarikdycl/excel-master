from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}


def inspect_workbook(file_path: str) -> dict:
    """
    Inspect an Excel workbook and return a structured report.

    The report contains:
    - workbook-level metadata
    - worksheet dimensions
    - data quality information
    - formulas
    - tables
    - charts
    - merged cells
    - filters
    - frozen panes
    - basic numeric statistics
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"Workbook not found: {path}")

    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {path.suffix}. "
            "Only .xlsx and .xlsm files are supported."
        )

    workbook = load_workbook(
        path,
        data_only=False
    )

    report = {
        "file": path.name,
        "path": str(path.resolve()),
        "sheet_count": len(workbook.sheetnames),
        "sheet_names": workbook.sheetnames,
        "sheets": []
    }

    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]

        try:
            df = pd.read_excel(
                path,
                sheet_name=sheet_name
            )
        except Exception:
            df = pd.DataFrame()

        non_empty_cells = 0
        formula_count = 0

        for row in worksheet.iter_rows():
            for cell in row:

                if cell.value is not None:
                    non_empty_cells += 1

                if (
                    isinstance(cell.value, str)
                    and cell.value.startswith("=")
                ):
                    formula_count += 1

        total_cells = (
            worksheet.max_row
            * worksheet.max_column
        )

        empty_cells = max(
            total_cells - non_empty_cells,
            0
        )

        merged_cells = [
            str(cell_range)
            for cell_range
            in worksheet.merged_cells.ranges
        ]

        tables = []

        for table_name, table in worksheet.tables.items():
            tables.append(
                {
                    "name": table_name,
                    "range": table.ref
                }
            )

        freeze_panes = None

        if worksheet.freeze_panes:
            freeze_panes = str(
                worksheet.freeze_panes
            )

        auto_filter = (
            worksheet.auto_filter.ref
            if worksheet.auto_filter.ref
            else None
        )

        chart_count = len(
            getattr(
                worksheet,
                "_charts",
                []
            )
        )

        image_count = len(
            getattr(
                worksheet,
                "_images",
                []
            )
        )

        sheet_report = {
            "name": sheet_name,

            # Excel structure
            "sheet_state": worksheet.sheet_state,
            "excel_rows": worksheet.max_row,
            "excel_columns": worksheet.max_column,
            "used_range": worksheet.calculate_dimension(),

            # Cell information
            "total_cells": total_cells,
            "non_empty_cells": non_empty_cells,
            "empty_cells": empty_cells,

            # Pandas / dataset information
            "data_rows": len(df),
            "columns": [
                str(column)
                for column in df.columns
            ],
            "data_types": {},
            "missing_values": {},
            "duplicate_rows": 0,
            "numeric_summary": {},

            # Excel features
            "formula_count": formula_count,
            "merged_cells": merged_cells,
            "merged_cell_count": len(
                merged_cells
            ),
            "tables": tables,
            "table_count": len(tables),
            "chart_count": chart_count,
            "image_count": image_count,
            "freeze_panes": freeze_panes,
            "auto_filter": auto_filter
        }

        if not df.empty:

            sheet_report["data_types"] = {
                str(column): str(dtype)
                for column, dtype
                in df.dtypes.items()
            }

            sheet_report["missing_values"] = {
                str(column): int(count)
                for column, count
                in df.isna().sum().items()
                if count > 0
            }

            sheet_report["duplicate_rows"] = int(
                df.duplicated().sum()
            )

            numeric_df = df.select_dtypes(
                include="number"
            )

            if not numeric_df.empty:

                sheet_report["numeric_summary"] = (
                    numeric_df
                    .describe()
                    .round(2)
                    .to_dict()
                )

        report["sheets"].append(
            sheet_report
        )

    return report