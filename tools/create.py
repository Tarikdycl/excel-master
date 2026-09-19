from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils import get_column_letter


DEFAULT_ROWS = 200


def create_workbook(spec: dict, output_path: str) -> str:
    """
    Generic Excel workbook builder.

    Expected spec format:

    {
        "title": "Inventory Tracker",
        "sheets": [
            {
                "name": "Inventory",
                "columns": [
                    {"name": "Product", "type": "text"},
                    {"name": "Quantity", "type": "number"},
                    {"name": "Price", "type": "currency"},
                    {
                        "name": "Status",
                        "type": "dropdown",
                        "options": ["In Stock", "Low Stock", "Out"]
                    }
                ]
            }
        ]
    }
    """

    validate_spec(spec)

    output = Path(output_path)
    output.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    wb = Workbook()

    default_sheet = wb.active
    wb.remove(default_sheet)

    for sheet_spec in spec["sheets"]:
        build_sheet(
            wb=wb,
            sheet_spec=sheet_spec
        )

    wb.save(output)

    return str(output)


def build_sheet(wb, sheet_spec: dict):
    sheet_name = sheet_spec["name"]

    ws = wb.create_sheet(
        title=sheet_name
    )

    columns = sheet_spec.get(
        "columns",
        []
    )

    # Optional title
    title = sheet_spec.get(
        "title"
    )

    header_row = 1

    if title:
        last_col = max(
            len(columns),
            1
        )

        end_column = get_column_letter(
            last_col
        )

        ws.merge_cells(
            f"A1:{end_column}1"
        )

        ws["A1"] = title

        ws["A1"].font = Font(
            size=16,
            bold=True,
            color="FFFFFF"
        )

        ws["A1"].fill = PatternFill(
            fill_type="solid",
            fgColor="16324F"
        )

        ws["A1"].alignment = Alignment(
            vertical="center"
        )

        ws.row_dimensions[1].height = 28

        header_row = 3

    # -------------------------
    # Headers
    # -------------------------

    headers = [
        column["name"]
        for column in columns
    ]

    if headers:
        for index, header in enumerate(
            headers,
            start=1
        ):
            cell = ws.cell(
                row=header_row,
                column=index
            )

            cell.value = header

        style_headers(
            ws,
            header_row,
            len(headers)
        )

    # -------------------------
    # Empty rows
    # -------------------------

    data_start_row = header_row + 1
    data_end_row = (
        data_start_row
        + sheet_spec.get(
            "rows",
            DEFAULT_ROWS
        )
        - 1
    )

    for row in range(
        data_start_row,
        data_end_row + 1
    ):
        for column_index in range(
            1,
            len(columns) + 1
        ):
            ws.cell(
                row=row,
                column=column_index
            )

    # -------------------------
    # Apply column rules
    # -------------------------

    for index, column_spec in enumerate(
        columns,
        start=1
    ):
        apply_column_type(
            ws=ws,
            column_index=index,
            column_spec=column_spec,
            start_row=data_start_row,
            end_row=data_end_row
        )

        apply_column_width(
            ws=ws,
            column_index=index,
            column_spec=column_spec
        )

        apply_formula(
            ws=ws,
            column_index=index,
            column_spec=column_spec,
            start_row=data_start_row,
            end_row=data_end_row,
            columns=columns
        )

    # -------------------------
    # Table
    # -------------------------

    if headers:
        last_column = get_column_letter(
            len(headers)
        )

        table_ref = (
            f"A{header_row}:"
            f"{last_column}{data_end_row}"
        )

        table_name = safe_table_name(
            sheet_name
        )

        table = Table(
            displayName=table_name,
            ref=table_ref
        )

        table.tableStyleInfo = (
            TableStyleInfo(
                name=sheet_spec.get(
                    "table_style",
                    "TableStyleMedium2"
                ),
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False
            )
        )

        ws.add_table(table)

    # -------------------------
    # Freeze panes
    # -------------------------

    if sheet_spec.get(
        "freeze_header",
        True
    ):
        ws.freeze_panes = (
            f"A{data_start_row}"
        )

    # -------------------------
    # Conditional formatting
    # -------------------------

    apply_conditional_rules(
        ws=ws,
        rules=sheet_spec.get(
            "conditional_rules",
            []
        ),
        columns=columns,
        start_row=data_start_row,
        end_row=data_end_row
    )

    # -------------------------
    # Charts
    # -------------------------

    apply_charts(
        ws=ws,
        charts=sheet_spec.get(
            "charts",
            []
        ),
        columns=columns,
        header_row=header_row,
        data_start_row=data_start_row,
        data_end_row=data_end_row
    )


def style_headers(
    ws,
    header_row,
    column_count
):
    fill = PatternFill(
        fill_type="solid",
        fgColor="2F75B5"
    )

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="FFFFFF"
        )
    )

    for column_index in range(
        1,
        column_count + 1
    ):
        cell = ws.cell(
            row=header_row,
            column=column_index
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = thin_border

    ws.row_dimensions[
        header_row
    ].height = 28


def apply_column_type(
    ws,
    column_index,
    column_spec,
    start_row,
    end_row
):
    column_letter = get_column_letter(
        column_index
    )

    column_type = column_spec.get(
        "type",
        "text"
    )

    for row in range(
        start_row,
        end_row + 1
    ):
        cell = ws[
            f"{column_letter}{row}"
        ]

        if column_type == "currency":
            cell.number_format = (
                '$#,##0.00'
            )

        elif column_type == "date":
            cell.number_format = (
                "mm/dd/yyyy"
            )

        elif column_type == "percent":
            cell.number_format = (
                "0.00%"
            )

        elif column_type == "integer":
            cell.number_format = "0"

        elif column_type == "number":
            cell.number_format = (
                "#,##0.00"
            )

    if column_type == "dropdown":
        options = column_spec.get(
            "options",
            []
        )

        if not options:
            return

        formula = (
            '"'
            + ",".join(
                str(option)
                for option in options
            )
            + '"'
        )

        validation = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True
        )

        ws.add_data_validation(
            validation
        )

        validation.add(
            f"{column_letter}"
            f"{start_row}:"
            f"{column_letter}"
            f"{end_row}"
        )


def apply_column_width(
    ws,
    column_index,
    column_spec
):
    column_letter = (
        get_column_letter(
            column_index
        )
    )

    width = column_spec.get(
        "width"
    )

    if width is None:
        name_length = len(
            column_spec.get(
                "name",
                ""
            )
        )

        width = max(
            12,
            min(
                name_length + 4,
                30
            )
        )

    ws.column_dimensions[
        column_letter
    ].width = width


def apply_formula(
    ws,
    column_index,
    column_spec,
    start_row,
    end_row,
    columns
):
    formula = column_spec.get(
        "formula"
    )

    if not formula:
        return

    column_letter = (
        get_column_letter(
            column_index
        )
    )

    for row in range(
        start_row,
        end_row + 1
    ):
        resolved_formula = (
            resolve_formula(
                formula=formula,
                row=row,
                columns=columns
            )
        )

        ws[
            f"{column_letter}{row}"
        ] = resolved_formula


def resolve_formula(
    formula: str,
    row: int,
    columns: list
):
    """
    Supports placeholders:

    {row}

    or column names:

    {Quantity}
    {Price}

    Example:
    =IF({Product}="","",{Quantity}*{Price})
    """

    result = formula.replace(
        "{row}",
        str(row)
    )

    for index, column in enumerate(
        columns,
        start=1
    ):
        column_name = column["name"]

        column_letter = (
            get_column_letter(
                index
            )
        )

        result = result.replace(
            f"{{{column_name}}}",
            f"{column_letter}{row}"
        )

    return result


def apply_conditional_rules(
    ws,
    rules,
    columns,
    start_row,
    end_row
):
    column_lookup = {
        column["name"]: (
            get_column_letter(
                index
            )
        )
        for index, column in enumerate(
            columns,
            start=1
        )
    }

    for rule in rules:
        column_name = rule.get(
            "column"
        )

        if column_name not in column_lookup:
            continue

        column_letter = (
            column_lookup[
                column_name
            ]
        )

        cell_range = (
            f"{column_letter}"
            f"{start_row}:"
            f"{column_letter}"
            f"{end_row}"
        )

        style = get_style(
            rule.get(
                "style",
                "warning"
            )
        )

        operator = rule.get(
            "operator"
        )

        value = rule.get(
            "value"
        )

        if operator == "greater_than":
            ws.conditional_formatting.add(
                cell_range,
                CellIsRule(
                    operator="greaterThan",
                    formula=[
                        str(value)
                    ],
                    fill=style["fill"],
                    font=style["font"]
                )
            )

        elif operator == "less_than":
            ws.conditional_formatting.add(
                cell_range,
                CellIsRule(
                    operator="lessThan",
                    formula=[
                        str(value)
                    ],
                    fill=style["fill"],
                    font=style["font"]
                )
            )

        elif operator == "equal":
            ws.conditional_formatting.add(
                cell_range,
                FormulaRule(
                    formula=[
                        f'{column_letter}'
                        f'{start_row}='
                        f'"{value}"'
                    ],
                    fill=style["fill"],
                    font=style["font"]
                )
            )

        elif operator == "not_equal":
            ws.conditional_formatting.add(
                cell_range,
                FormulaRule(
                    formula=[
                        f'{column_letter}'
                        f'{start_row}<>'
                        f'"{value}"'
                    ],
                    fill=style["fill"],
                    font=style["font"]
                )
            )


def get_style(style_name: str):
    styles = {
        "danger": {
            "fill": PatternFill(
                fill_type="solid",
                fgColor="F4CCCC"
            ),
            "font": Font(
                color="990000"
            )
        },

        "warning": {
            "fill": PatternFill(
                fill_type="solid",
                fgColor="FFF2CC"
            ),
            "font": Font(
                color="7F6000"
            )
        },

        "success": {
            "fill": PatternFill(
                fill_type="solid",
                fgColor="D9EAD3"
            ),
            "font": Font(
                color="274E13"
            )
        },

        "info": {
            "fill": PatternFill(
                fill_type="solid",
                fgColor="D9EAF7"
            ),
            "font": Font(
                color="16324F"
            )
        }
    }

    return styles.get(
        style_name,
        styles["warning"]
    )


def apply_charts(
    ws,
    charts,
    columns,
    header_row,
    data_start_row,
    data_end_row
):
    column_lookup = {
        column["name"]: index
        for index, column in enumerate(
            columns,
            start=1
        )
    }

    for chart_spec in charts:
        chart_type = chart_spec.get(
            "type",
            "bar"
        )

        data_column = chart_spec.get(
            "data_column"
        )

        category_column = (
            chart_spec.get(
                "category_column"
            )
        )

        if (
            data_column
            not in column_lookup
        ):
            continue

        if (
            category_column
            not in column_lookup
        ):
            continue

        if chart_type == "line":
            chart = LineChart()

        elif chart_type == "pie":
            chart = PieChart()

        else:
            chart = BarChart()

        chart.title = chart_spec.get(
            "title",
            "Chart"
        )

        data = Reference(
            ws,
            min_col=column_lookup[
                data_column
            ],
            min_row=header_row,
            max_row=data_end_row
        )

        categories = Reference(
            ws,
            min_col=column_lookup[
                category_column
            ],
            min_row=data_start_row,
            max_row=data_end_row
        )

        chart.add_data(
            data,
            titles_from_data=True
        )

        chart.set_categories(
            categories
        )

        position = chart_spec.get(
            "position",
            "H2"
        )

        ws.add_chart(
            chart,
            position
        )


def safe_table_name(
    name: str
) -> str:
    cleaned = re.sub(
        r"[^A-Za-z0-9_]",
        "",
        name
    )

    if not cleaned:
        cleaned = "Sheet"

    if cleaned[0].isdigit():
        cleaned = (
            "Table"
            + cleaned
        )

    return (
        cleaned
        + "Table"
    )


def validate_spec(
    spec: dict
):
    if not isinstance(
        spec,
        dict
    ):
        raise ValueError(
            "Spec must be a dictionary."
        )

    sheets = spec.get(
        "sheets"
    )

    if not sheets:
        raise ValueError(
            "Spec must contain at least one sheet."
        )

    for sheet in sheets:
        if "name" not in sheet:
            raise ValueError(
                "Every sheet must have a name."
            )

        columns = sheet.get(
            "columns",
            []
        )

        for column in columns:
            if "name" not in column:
                raise ValueError(
                    "Every column must have a name."
                )